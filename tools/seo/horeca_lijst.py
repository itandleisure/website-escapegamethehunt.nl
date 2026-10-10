"""Maakt per stad een lijst met horeca rond het speelveld, voor acquisitie van horecapartners.

Gebruik:  python tools/seo/horeca_lijst.py groningen [utrecht ...]
          python tools/seo/horeca_lijst.py alle

Zoekt via DataForSEO (Google Maps) naar restaurants, cafés en bars rond de startlocatie,
en schrijft tools/seo/horeca/horeca-<stad>.xlsx met naam, soort, adres, telefoon, website,
Google-score, aantal reviews, afstand tot de start en of het binnen het speelveld ligt.
Nodig: DATAFORSEO_LOGIN en DATAFORSEO_PASSWORD. Kost ongeveer $0,01 per stad.
"""
import base64, json, math, os, sys, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
API = 'https://api.dataforseo.com/v3/serp/google/maps/live/advanced'
QUERIES = ['restaurant', 'café', 'bar', 'lunchroom', 'grand café']


def search(auth, kw, lat, lon):
    body = json.dumps([{'keyword': kw, 'location_coordinate': f'{lat},{lon},17', 'language_code': 'nl', 'depth': 60}]).encode()
    req = urllib.request.Request(API, data=body, headers={'Authorization': 'Basic ' + auth, 'Content-Type': 'application/json'})
    r = json.load(urllib.request.urlopen(req, timeout=300))
    return (r['tasks'][0].get('result') or [{}])[0].get('items') or [], r.get('cost') or 0


def inside(lat, lon, poly):
    c = False
    for (a1, b1), (a2, b2) in zip(poly, poly[1:] + poly[:1]):
        if (b1 > lon) != (b2 > lon) and lat < (a2 - a1) * (lon - b1) / (b2 - b1) + a1:
            c = not c
    return c


def dist(a, b):
    return math.hypot((a[0] - b[0]) * 111000, (a[1] - b[1]) * 111000 * math.cos(math.radians(a[0])))


def city(slug, auth):
    cfg_f = os.path.join(ROOT, 'tools', 'kaarten', f'speelveld-{slug}.json')
    if not os.path.exists(cfg_f):
        print(slug, ': nog geen speelveld, overgeslagen'); return 0
    cfg = json.load(open(cfg_f, encoding='utf-8'))
    poly = [tuple(p) for p in cfg['field']]
    lat, lon = (cfg['start'][1], cfg['start'][2]) if cfg.get('start') else (sum(p[0] for p in poly) / len(poly), sum(p[1] for p in poly) / len(poly))
    seen, rows, cost = set(), [], 0
    for q in QUERIES:
        items, c = search(auth, q, lat, lon); cost += c
        for i in items:
            key = i.get('place_id') or i.get('title')
            if key in seen or not i.get('latitude'):
                continue
            seen.add(key)
            pos = (i['latitude'], i['longitude'])
            rt = i.get('rating') or {}
            rows.append(dict(naam=i.get('title'), soort=i.get('category') or '', adres=i.get('address') or '',
                             telefoon=i.get('phone') or '', website=i.get('url') or '',
                             score=rt.get('value'), reviews=rt.get('votes_count') or 0,
                             afstand=round(dist(pos, (lat, lon))), in_speelveld='ja' if inside(pos[0], pos[1], poly) else 'nee',
                             maps=i.get('check_url') or ''))
    rows = [r for r in rows if r['in_speelveld'] == 'ja' or r['afstand'] < 600]
    for r in rows:  # simpele prioriteit: goede score, veel reviews, dicht bij de start
        r['prioriteit'] = round((r['score'] or 0) * math.log10(1 + r['reviews']) - r['afstand'] / 500, 2)
    rows.sort(key=lambda r: -r['prioriteit'])
    wb = Workbook(); ws = wb.active; ws.title = slug.capitalize()
    cols = ['prioriteit', 'naam', 'soort', 'adres', 'telefoon', 'website', 'score', 'reviews', 'afstand', 'in_speelveld',
            'benaderd', 'contactpersoon', 'status', 'notities']
    heads = ['Prioriteit', 'Naam', 'Soort', 'Adres', 'Telefoon', 'Website', 'Google-score', 'Reviews', 'Afstand tot start (m)',
             'In speelveld', 'Benaderd op', 'Contactpersoon', 'Status', 'Notities']
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='29394A')
    for r in rows:
        ws.append([r.get(k, '') for k in cols])
    for i, w in enumerate([10, 32, 20, 38, 16, 34, 12, 10, 12, 12, 14, 20, 14, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.join(HERE, 'horeca'), exist_ok=True)
    out = os.path.join(HERE, 'horeca', f'horeca-{slug}.xlsx')
    wb.save(out)
    print(f'{slug}: {len(rows)} zaken, kosten ${cost:.3f} -> {out}')
    return cost


if __name__ == '__main__':
    auth = base64.b64encode(f"{os.environ['DATAFORSEO_LOGIN']}:{os.environ['DATAFORSEO_PASSWORD']}".encode()).decode()
    slugs = sys.argv[1:]
    if slugs == ['alle']:
        slugs = sorted(f[len('speelveld-'):-5] for f in os.listdir(os.path.join(ROOT, 'tools', 'kaarten')) if f.startswith('speelveld-') and f.endswith('.json'))
    total = sum(city(s, auth) for s in slugs)
    print(f'totaal ${total:.2f}')
