"""Zet outreach-mails als concept in de eigen mailbox (IMAP), zodat je ze alleen nog hoeft na te lezen en te versturen.

Gebruik:
  python tools/seo/outreach/concepten.py            concepten klaarzetten voor alle regels met status "Klaar"
  python tools/seo/outreach/concepten.py --test     niets naar de mailbox; schrijft .eml-bestanden in outreach/test/
  python tools/seo/outreach/concepten.py --mappen   toont de mappen in de mailbox (om de conceptenmap te vinden)

Leest het tabblad "Mails" van tools/seo/outreach/outreach.xlsx (kolommen Stad, Site, Aan, Onderwerp, Tekst, Status).
Na het klaarzetten wordt de status "Concept klaargezet <datum>", zodat een regel nooit twee keer in je mailbox komt.
Er wordt niets verstuurd: dat doe je zelf vanuit je mailprogramma.

Inloggegevens staan in tools/seo/outreach/outreach.local.json (niet in git), zie outreach.local.example.json.
"""
import argparse, datetime as dt, imaplib, json, os, re, sys, time
from email.message import EmailMessage
from email.utils import formatdate, make_msgid
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'outreach.xlsx')
CFG = os.path.join(HERE, 'outreach.local.json')


def kolommen(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and 'Aan' in row and 'Status' in row:
            return i, {str(v): j for j, v in enumerate(row) if v}
    sys.exit('Kopregel met "Aan" en "Status" niet gevonden in tabblad Mails.')


def bericht(cfg, aan, onderwerp, tekst):
    m = EmailMessage()
    m['From'] = cfg['van']
    m['To'] = aan
    m['Subject'] = onderwerp
    m['Date'] = formatdate(localtime=True)
    m['Message-ID'] = make_msgid(domain=cfg['van'].split('@')[-1].strip('> '))
    if cfg.get('bcc'):
        m['Bcc'] = cfg['bcc']
    m.set_content(tekst)
    return m


def conceptmap(imap, cfg):
    if cfg.get('concepten_map'):
        return cfg['concepten_map']
    typ, data = imap.list()
    namen = []
    for regel in data or []:
        s = regel.decode(errors='replace')
        naam = re.search(r'"([^"]*)"\s*$|(\S+)\s*$', s)
        naam = naam.group(1) or naam.group(2)
        if '\\Drafts' in s:
            return naam
        namen.append(naam)
    for kand in ('Drafts', 'INBOX.Drafts', 'Concepten', 'INBOX.Concepten', 'Concepts'):
        if kand in namen:
            return kand
    sys.exit('Conceptenmap niet gevonden. Draai met --mappen en zet "concepten_map" in outreach.local.json.')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--test', action='store_true')
    ap.add_argument('--mappen', action='store_true')
    a = ap.parse_args()
    if not os.path.exists(CFG) and not a.test:
        sys.exit(f'Maak eerst {CFG} aan (zie outreach.local.example.json).')
    cfg = json.load(open(CFG, encoding='utf-8')) if os.path.exists(CFG) else {'van': 'Escape Game The Hunt <info@escapegamethehunt.nl>'}

    imap = None
    if not a.test:
        imap = imaplib.IMAP4_SSL(cfg['imap_host'], int(cfg.get('imap_poort', 993)))
        imap.login(cfg['gebruiker'], cfg['wachtwoord'])
        if a.mappen:
            for regel in imap.list()[1]:
                print(regel.decode(errors='replace'))
            imap.logout(); return
        doel = conceptmap(imap, cfg)
        print('Conceptenmap:', doel)

    wb = load_workbook(XLSX)
    ws = wb['Mails']
    kop, k = kolommen(ws)
    n = 0
    for r in range(kop + 1, ws.max_row + 1):
        cel = lambda naam: ws.cell(r, k[naam] + 1)
        status = str(cel('Status').value or '').strip().lower()
        aan = str(cel('Aan').value or '').strip()
        if status != 'klaar' or '@' not in aan:
            continue
        m = bericht(cfg, aan, str(cel('Onderwerp').value), str(cel('Tekst').value))
        if a.test:
            os.makedirs(os.path.join(HERE, 'test'), exist_ok=True)
            f = os.path.join(HERE, 'test', re.sub(r'[^\w.-]', '_', str(cel('Site').value)) + '.eml')
            open(f, 'wb').write(bytes(m))
            print('test:', f)
        else:
            typ, resp = imap.append(doel, '(\\Draft)', imaplib.Time2Internaldate(time.time()), bytes(m))
            if typ != 'OK':
                print('MISLUKT:', cel('Site').value, resp); continue
            cel('Status').value = 'Concept klaargezet ' + dt.date.today().strftime('%d-%m-%Y')
            print('concept:', cel('Site').value, '->', aan)
        n += 1
    if imap:
        imap.logout()
        wb.save(XLSX)
    print(f'{n} concept(en) {"getest" if a.test else "in je mailbox gezet"}.')


if __name__ == '__main__':
    main()
