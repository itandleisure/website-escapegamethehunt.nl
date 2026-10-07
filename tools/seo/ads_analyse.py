"""Google Ads-analyse: kosten per zoekwoord naast onze organische positie, met advies.

Gebruik:  python tools/seo/ads_analyse.py
Leest de zoekvolumes/klikprijzen uit tools/seo/ads/google-ads-kosten-per-stad.xlsx (tabblad "Per stad"),
meet per zoekwoord de huidige positie in Google (mobiel, vanuit het centrum van de stad, top 30, zoals de
weekmeting) en zet de Search Console-cijfers van de laatste 28 dagen erbij. Schrijft het tabblad "Analyse".
"""
import base64, concurrent.futures as cf, glob, json, os, urllib.request
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
F = os.path.join(HERE, 'ads', 'google-ads-kosten-per-stad.xlsx')
SITE = 'escapegamethehunt.nl'
LOCS = {l['name']: l for l in json.load(open(os.path.join(ROOT, 'src', 'data', 'locations.json'), encoding='utf-8'))}
AUTH = base64.b64encode(f"{os.environ['DATAFORSEO_LOGIN']}:{os.environ['DATAFORSEO_PASSWORD']}".encode()).decode()


def positie(stad, kw):
    l = LOCS[stad]
    body = json.dumps([{'keyword': kw, 'language_code': 'nl', 'location_coordinate': f"{l['lat']:.5f},{l['lng']:.5f},2000",
                        'device': 'mobile', 'os': 'android', 'depth': 30}]).encode()
    req = urllib.request.Request('https://api.dataforseo.com/v3/serp/google/organic/live/advanced', data=body,
                                 headers={'Authorization': 'Basic ' + AUTH, 'Content-Type': 'application/json'})
    try:
        r = json.load(urllib.request.urlopen(req, timeout=300))
        t = r['tasks'][0]
        if t.get('status_code') != 20000:
            return 'ERR'
        items = (t.get('result') or [{}])[0].get('items') or []
        if sum(1 for i in items if i.get('type') == 'organic') < 15:
            return 'ERR'  # onvolledig resultaat (komt voor bij drukte)
    except Exception:
        return 'ERR'  # niet als "niet gevonden" opslaan, later opnieuw proberen
    ads = any(i.get('type') == 'paid' for i in items)
    for i in items:
        if i.get('type') == 'organic' and SITE in (i.get('domain') or ''):
            return i['rank_group'], i['url'].replace('https://' + SITE, ''), ads
    return None, '', ads


def gsc():
    f = sorted(glob.glob(os.path.join(HERE, 'data', 'gsc-*.json')))
    if not f:
        return {}
    return {r['keys'][0]: r for r in json.load(open(f[-1], encoding='utf-8'))['q_nu']}


def advies(pos, vol, cpc, kw, gp):
    merk = kw.startswith('the hunt')
    if merk:
        return ('Niet nodig: organisch al bovenaan' if pos and pos <= 3 else 'Wel doen: goedkoop, eigen naam moet bovenaan'), 'merk'
    if pos and pos <= 3:
        return 'Niet adverteren: organisch al top 3', 'nee'
    if kw.startswith('escape room'):
        return 'Liever niet: zoekers willen vaak een binnen-escaperoom', 'nee'
    if vol < 30:
        return 'Te weinig zoekvolume', 'nee'
    if pos and pos <= 10:
        return f'Testen: organisch {pos}, advertentie bovenaan pakt extra klikken', 'test'
    return 'Adverteren zinvol: organisch niet zichtbaar' + (' (gemiddeld ' + f'{gp:.0f}' + ' in Search Console)' if gp else ''), 'ja'


def main():
    wb = load_workbook(F)
    src = wb['Per stad']
    rows = [r for r in src.iter_rows(min_row=6, values_only=True) if r[0]]
    cache_f = os.path.join(HERE, 'ads', 'posities-cache.json')
    cache = json.load(open(cache_f, encoding='utf-8')) if os.path.exists(cache_f) else {}
    for poging in range(4):
        todo = [r for r in rows if f'{r[0]}|{r[1]}' not in cache]
        if not todo:
            break
        with cf.ThreadPoolExecutor(4) as ex:
            for r, p in zip(todo, ex.map(lambda r: positie(r[0], r[1]), todo)):
                if p != 'ERR':
                    cache[f'{r[0]}|{r[1]}'] = p
        print(f'poging {poging + 1}: {len(todo)} gemeten, nog {sum(1 for r in rows if f"{r[0]}|{r[1]}" not in cache)} open')
    json.dump(cache, open(cache_f, 'w', encoding='utf-8'), ensure_ascii=False)
    pos = [tuple(cache.get(f'{r[0]}|{r[1]}', (None, 'niet gemeten', False))) for r in rows]
    G = gsc()
    if 'Analyse' in wb.sheetnames:
        del wb['Analyse']
    ws = wb.create_sheet('Analyse', 0)
    intro = ['Adverteren of niet? Kosten per zoekwoord naast onze positie in Google (gemeten 7-10-2026, mobiel, vanuit het centrum van de stad).',
             'Positie = organische plek nu (leeg = niet in de top 30). Search Console = gemiddelde positie, vertoningen en klikken van de laatste 28 dagen.',
             'Advies: top 3 organisch = niet adverteren (je krijgt de klik al gratis). Positie 4-10 = testen. Niet zichtbaar + genoeg volume = adverteren zinvol.',
             'Bod bovenaan laag/hoog = de bandbreedte die Google Ads geeft voor een plek bovenaan de pagina (Zoekwoordplanner, gemiddelde van 12 maanden, heel NL). Geen live prijzen: de echte klikprijs hangt af van de veiling en de kwaliteitsscore.',
             '"Er adverteren al anderen" = er stonden betaalde advertenties boven de zoekresultaten: daar is concurrentie, dus de klikprijs kan hoger uitvallen.']
    for i, t in enumerate(intro, 1):
        ws.cell(i, 1, t).font = Font(bold=i == 1, italic=i > 1)
    BOD = json.load(open(os.path.join(HERE, 'ads', 'zoekwoorden-ads.json'), encoding='utf-8'))
    H = ['Stad', 'Zoekwoord', 'Zoekvolume/maand', 'Klikprijs gem. (€)', 'Bod bovenaan laag (€)', 'Bod bovenaan hoog (€)', 'Geschatte kosten/maand (€)', 'Kosten/maand laag-hoog (€)', 'Onze positie nu', 'Onze pagina',
         'Search Console gem. positie', 'SC vertoningen (28 d)', 'SC klikken (28 d)', 'Er adverteren al anderen', 'Advies']
    top = len(intro) + 2
    for c, h in enumerate(H, 1):
        x = ws.cell(top, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
        x.alignment = Alignment(wrap_text=True, vertical='top')
    kleur = {'ja': 'D9F2DC', 'test': 'FFF3C4', 'nee': 'F3F5F7', 'merk': 'E3ECF7'}
    samen = {}
    for i, (r, (p, url, ads)) in enumerate(zip(rows, pos), top + 1):
        g = G.get(r[1], {})
        gp = g.get('position')
        adv, k = advies(p, r[2], r[3], r[1], gp)
        b = BOD.get(r[1], {}); lo, hi = b.get('low'), b.get('high')
        band = f'{r[5] * lo:.0f} - {r[5] * hi:.0f}' if lo and hi else ''
        vals = [r[0], r[1], r[2], r[3], lo or '', hi or '', r[6], band, p or '–', url, round(gp, 1) if gp else '', g.get('impressions', ''), g.get('clicks', ''), 'ja' if ads else 'nee', adv]
        for c, v in enumerate(vals, 1):
            x = ws.cell(i, c, v); x.fill = PatternFill('solid', fgColor=kleur[k]); x.alignment = Alignment(vertical='top', wrap_text=c in (10, 15))
        s = samen.setdefault(k, [0, 0.0, 0]); s[0] += 1; s[1] += r[6] or 0; s[2] += r[2]
    for col, w in zip('ABCDEFGHIJKLMNO', (13, 30, 12, 10, 10, 10, 13, 14, 10, 38, 12, 12, 10, 12, 52)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(top + 1, 3); ws.auto_filter.ref = f'A{top}:O{top + len(rows)}'
    try:
        wb.save(F); out = F
    except PermissionError:  # bestand staat open in Excel
        out = F.replace('.xlsx', '-analyse.xlsx'); wb.save(out)
    print(out)
    for k, (n, kost, vol) in samen.items():
        print(k, n, 'zoekwoorden', vol, 'zoekopdr.', round(kost), 'euro/mnd')
    return rows, pos


if __name__ == '__main__':
    main()
