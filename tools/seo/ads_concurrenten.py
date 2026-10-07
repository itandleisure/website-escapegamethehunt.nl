"""Welke concurrenten adverteren op onze zoekwoorden?

Gebruik:  python tools/seo/ads_concurrenten.py
Haalt voor elk zoekwoord uit tools/seo/ads/zoekwoorden-ads.json (plus extra zoekwoorden voor de prioriteitssteden)
de live Google-resultaten op (mobiel en desktop, vanuit het centrum van de stad) en noteert elke advertentie:
adverteerder (domein), plek, titel en tekst. Daarnaast per bekende concurrent de betaalde zoekwoorden uit de
DataForSEO-database. Resultaat: tools/seo/ads/concurrenten-advertenties.xlsx (+ cache in concurrenten-cache.json).
"""
import base64, collections, concurrent.futures as cf, json, os, re, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
ADS = os.path.join(HERE, 'ads')
AUTH = base64.b64encode(f"{os.environ['DATAFORSEO_LOGIN']}:{os.environ['DATAFORSEO_PASSWORD']}".encode()).decode()
LOCS = {l['name']: l for l in json.load(open(os.path.join(ROOT, 'src', 'data', 'locations.json'), encoding='utf-8')) if l.get('start')}
PRIO = ['Utrecht', 'Groningen', 'Arnhem', 'Nijmegen', 'Rotterdam', 'Leeuwarden']
EXTRA = ['uitje {s}', 'groepsuitje {s}', 'personeelsuitje {s}', 'citygame {s}', 'escape game {s}', 'outdoor escape room {s}']
CONCURRENTEN = ['tbevents.nl', '1001activiteiten.nl', 'uitjesbazen.nl', 'onemotion.nl', 'cityadventures.nl', 'mycityhunt.nl', 'outsideescape.nl',
                'quest-escapes.nl', 'stadsarrangement.nl', 'neverrest.nl', 'escapehunt.com', 'uitjesbureau.nl', 'coddygames.com', 'getevents.nl',
                'citygame.com', 'monsterevents.nl', 'uitjeseneten.nl', 'crazy88.nl', 'escapetours.com', 'funvillage.nl', 'pubevents.nl']


def post(path, task):
    req = urllib.request.Request('https://api.dataforseo.com/v3/' + path, data=json.dumps([task]).encode(),
                                 headers={'Authorization': 'Basic ' + AUTH, 'Content-Type': 'application/json'})
    return json.load(urllib.request.urlopen(req, timeout=300))['tasks'][0]


def serp(job):
    stad, kw, device = job
    l = LOCS[stad]
    try:
        t = post('serp/google/organic/live/advanced', {'keyword': kw, 'language_code': 'nl', 'device': device,
                 'location_coordinate': f"{l['lat']:.5f},{l['lng']:.5f},2000", 'depth': 20})
        if t.get('status_code') != 20000:
            return 'ERR'
        items = (t.get('result') or [{}])[0].get('items') or []
        if sum(1 for i in items if i.get('type') == 'organic') < 8:
            return 'ERR'
    except Exception:
        return 'ERR'
    return [{'domein': (i.get('domain') or '').replace('www.', ''), 'plek': i.get('rank_group'), 'boven': (i.get('rank_absolute') or 99) < 5,
             'titel': i.get('title') or '', 'tekst': (i.get('description') or '')[:200]} for i in items if i.get('type') == 'paid']


def labs(domein):
    try:
        t = post('dataforseo_labs/google/ranked_keywords/live', {'target': domein, 'location_code': 2528, 'language_code': 'nl',
                 'item_types': ['paid'], 'limit': 100, 'order_by': ['keyword_data.keyword_info.search_volume,desc']})
        return [(it['keyword_data']['keyword'], it['keyword_data']['keyword_info'].get('search_volume') or 0,
                 it['ranked_serp_element']['serp_item'].get('title') or '') for it in ((t.get('result') or [{}])[0].get('items') or [])]
    except Exception:
        return []


def main():
    kws = json.load(open(os.path.join(ADS, 'zoekwoorden-ads.json'), encoding='utf-8'))
    jobs = []
    for stad in LOCS:
        for kw in kws:
            if kw.endswith(' ' + stad.lower()) and (kws[kw].get('vol') or 0) >= 20:
                jobs.append((stad, kw))
    for stad in PRIO:
        jobs += [(stad, p.format(s=stad.lower())) for p in EXTRA if (stad, p.format(s=stad.lower())) not in jobs]
    jobs = [(s, k, d) for s, k in jobs for d in ('mobile', 'desktop')]
    cache_f = os.path.join(ADS, 'concurrenten-cache.json')
    cache = json.load(open(cache_f, encoding='utf-8')) if os.path.exists(cache_f) else {}
    for poging in range(4):
        todo = [j for j in jobs if '|'.join(j) not in cache]
        if not todo:
            break
        with cf.ThreadPoolExecutor(4) as ex:
            for j, res in zip(todo, ex.map(serp, todo)):
                if res != 'ERR':
                    cache['|'.join(j)] = res
        json.dump(cache, open(cache_f, 'w', encoding='utf-8'), ensure_ascii=False)
        print(f'poging {poging + 1}: {len(todo)} gemeten, nog {sum(1 for j in jobs if "|".join(j) not in cache)} open')
    with cf.ThreadPoolExecutor(6) as ex:
        L = dict(zip(CONCURRENTEN, ex.map(labs, CONCURRENTEN)))

    wb = Workbook()
    # 1. per adverteerder
    per = collections.defaultdict(lambda: {'kw': set(), 'steden': set(), 'n': 0, 'titels': collections.Counter()})
    rows = []
    for key, ads in cache.items():
        stad, kw, dev = key.split('|')
        vol = (kws.get(kw) or {}).get('vol') or ''
        for a in ads:
            if 'escapegamethehunt' in a['domein']:
                continue
            p = per[a['domein']]; p['kw'].add(kw); p['steden'].add(stad); p['n'] += 1; p['titels'][a['titel']] += 1
            rows.append([stad, kw, vol, dev, a['domein'], a['plek'], a['titel'], a['tekst']])
    ws = wb.active; ws.title = 'Per adverteerder'
    intro = ['Wie adverteert er op onze zoekwoorden? Live gemeten in Google (7-10-2026), mobiel en desktop, vanuit het centrum van elke stad.',
             f'Gemeten: {len(jobs) // 2} zoekwoorden x 2 apparaten. "Keer gezien" = aantal zoekresultaten waarin de adverteerder een advertentie had.']
    for i, t in enumerate(intro, 1):
        ws.cell(i, 1, t).font = Font(bold=i == 1, italic=i > 1)
    H = ['Adverteerder', 'Keer gezien', 'Aantal zoekwoorden', 'Steden', 'Zoekwoorden', 'Meest gebruikte advertentietitel']
    top = 4
    for c, h in enumerate(H, 1):
        x = ws.cell(top, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
    for r, (d, p) in enumerate(sorted(per.items(), key=lambda x: -x[1]['n']), top + 1):
        for c, v in enumerate([d, p['n'], len(p['kw']), ', '.join(sorted(p['steden'])), ', '.join(sorted(p['kw'])), p['titels'].most_common(1)[0][0]], 1):
            ws.cell(r, c, v).alignment = Alignment(wrap_text=c in (4, 5, 6), vertical='top')
    for col, w in zip('ABCDEF', (28, 11, 11, 40, 70, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'B5'
    # 2. alle advertenties
    ws = wb.create_sheet('Alle advertenties')
    H = ['Stad', 'Zoekwoord', 'Zoekvolume/maand', 'Apparaat', 'Adverteerder', 'Plek', 'Advertentietitel', 'Advertentietekst']
    for c, h in enumerate(H, 1):
        x = ws.cell(1, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
    for r, row in enumerate(sorted(rows), 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v).alignment = Alignment(wrap_text=c in (7, 8), vertical='top')
    for col, w in zip('ABCDEFGH', (13, 30, 12, 10, 26, 7, 50, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:H{len(rows) + 1}'
    # 3. per zoekwoord
    ws = wb.create_sheet('Per zoekwoord')
    H = ['Stad', 'Zoekwoord', 'Zoekvolume/maand', 'Aantal adverteerders', 'Adverteerders']
    for c, h in enumerate(H, 1):
        x = ws.cell(1, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
    pk = collections.defaultdict(set)
    for key, ads in cache.items():
        stad, kw, dev = key.split('|')
        pk[(stad, kw)] |= {a['domein'] for a in ads if 'escapegamethehunt' not in a['domein']}
    for r, ((stad, kw), ds) in enumerate(sorted(pk.items(), key=lambda x: (x[0][0], -len(x[1]))), 2):
        for c, v in enumerate([stad, kw, (kws.get(kw) or {}).get('vol') or '', len(ds), ', '.join(sorted(ds))], 1):
            ws.cell(r, c, v)
    for col, w in zip('ABCDE', (13, 30, 12, 12, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:E{len(pk) + 1}'
    # 4. database
    ws = wb.create_sheet('Database (landelijk)')
    ws.cell(1, 1, 'Betaalde zoekwoorden per concurrent volgens de DataForSEO-database (landelijk, onvolledig: alleen wat hun crawler heeft gezien).').font = Font(italic=True)
    for c, h in enumerate(['Concurrent', 'Zoekwoord', 'Zoekvolume/maand', 'Advertentietitel'], 1):
        x = ws.cell(3, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
    r = 4
    for d, lst in L.items():
        for kw, vol, tit in lst:
            for c, v in enumerate([d, kw, vol, tit], 1):
                ws.cell(r, c, v)
            r += 1
    for col, w in zip('ABCD', (26, 40, 14, 60)):
        ws.column_dimensions[col].width = w
    f = os.path.join(ADS, 'concurrenten-advertenties.xlsx')
    try:
        wb.save(f)
    except PermissionError:
        f = f.replace('.xlsx', '-nieuw.xlsx'); wb.save(f)
    print(f)
    for d, p in sorted(per.items(), key=lambda x: -x[1]['n'])[:20]:
        print(f"{p['n']:>3} {d:28} {len(p['kw']):>3} zw  {', '.join(sorted(p['steden']))[:80]}")
    print({d: len(v) for d, v in L.items() if v})


if __name__ == '__main__':
    main()
