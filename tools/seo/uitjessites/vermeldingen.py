"""Kant-en-klare teksten voor vermeldingen op uitjessites, VVV-sites en Tripadvisor, per stad.

Gebruik:  python tools/seo/uitjessites/vermeldingen.py
Maakt tools/seo/uitjessites/vermeldingen-per-stad.xlsx: één regel per stad met een speelveld (start bekend),
plus een landelijke regel. Feiten komen uit src/data (start, prijs) en tools/kaarten (straten in het speelveld),
zodat de teksten gelijk blijven met de website.
"""
import json, os, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
SITE = json.load(open(os.path.join(ROOT, 'src', 'data', 'site.json'), encoding='utf-8'))
P = SITE['prices']
URL = 'https://escapegamethehunt.nl'
LOCS = [l for l in json.load(open(os.path.join(ROOT, 'src', 'data', 'locations.json'), encoding='utf-8')) if l.get('start')]

KORT = [
    'Outdoor escape room in {c}: los puzzels op in de binnenstad en ontsnap aan de Hunters. 90 minuten, vanaf 8 personen.',
    'Geen kamer maar heel {c} als speelveld: puzzelen, vluchten voor de Hunters en het extractiepunt vinden. Vanaf 8 personen.',
    'Net als in Hunted, maar dan zelf: in {c} lossen teams puzzels op terwijl Hunters jullie opjagen. Start: {s}.',
    'Een escape game buiten in {c} voor teams, vrienden en vrijgezellen: puzzels, Hunters en 90 minuten spanning.',
]
OPEN = [
    'Escape Game The Hunt is een outdoor escape room in het centrum van {c}. Jullie spelen niet in een kamer, maar in de straten van de binnenstad.',
    'Bij Escape Game The Hunt wordt de binnenstad van {c} één grote escape room. Jullie team is op de vlucht en heeft 90 minuten om te ontsnappen.',
    'Escape Game The Hunt brengt de spanning van een escape room naar buiten, midden in {c}. Geïnspireerd op tv-programma\'s als Hunted en Jachtseizoen.',
    'Zin in een uitje waarbij je samenwerkt, puzzelt en rent? Escape Game The Hunt speel je buiten in de binnenstad van {c}.',
]
HOE = ('Elk team krijgt een gametas met puzzels en een GameApp. Wie de puzzels oplost, vindt de GPS-code van het geheime extractiepunt. '
       'Ondertussen gaan de Hunters op jacht: zij zien elke 10 minuten waar jullie zijn. Binnen het speelveld mogen jullie je vrij bewegen, '
       'dus elk team kiest zijn eigen route en elk spel loopt anders.')
WIE = ('Geschikt voor teamuitjes, bedrijfsuitjes, vrijgezellenfeesten, verjaardagen en schoolgroepen. Vanaf {mn} personen, in teams van ongeveer zes; '
       'groepen tot ongeveer {mx} deelnemers zijn mogelijk. Het spel gaat ook door als het regent.')
PRIJS = '€{b} ex. btw voor de hele groep tot en met {bm} personen, daarboven €{pp} per persoon (ex. btw).'


def labels(slug):
    f = os.path.join(ROOT, 'tools', 'kaarten', f'speelveld-{slug}.json')
    if not os.path.exists(f):
        return []
    return [l[0] if isinstance(l, list) else l for l in json.load(open(f, encoding='utf-8')).get('labels', [])]


def lijst(xs):
    return ', '.join(xs[:-1]) + ' en ' + xs[-1] if len(xs) > 1 else (xs[0] if xs else '')


def rij(i, l):
    c, s = l['name'], l['start']
    lab = [x for x in labels(l['slug']) if x.lower() not in s.lower()][:4]
    veld = f'Het speelveld loopt door de binnenstad, langs onder meer {lijst(lab)}.' if lab else ''
    lang = ' '.join(x for x in [OPEN[i % 4].format(c=c), f'Startlocatie: {s}.', veld, HOE,
                                WIE.format(mn=P['min_people'], mx=P['max_people'])] if x)
    prijs = PRIJS.format(b=P['base'], bm=P['base_max_people'], pp=P['per_person'])
    foto = URL + (l.get('image') or SITE.get('og_image', ''))
    kaart = URL + l['field_map'] + '.jpg' if l.get('field_map') else ''
    return [f'Escape Game The Hunt {c}', KORT[i % 4].format(c=c, s=s), lang, f'{s}, {c}', URL + l['url'], prijs,
            '90 minuten (met uitleg ongeveer 2 uur)', f'{P["min_people"]} tot ca. {P["max_people"]} personen', foto, kaart]


def main():
    H = ['Naam', 'Korte omschrijving (max. 160 tekens)', 'Lange omschrijving', 'Startlocatie', 'Website (link naar)', 'Prijs',
         'Duur', 'Groepsgrootte', 'Foto', 'Kaart speelveld']
    rows = [rij(i, l) for i, l in enumerate(LOCS)]
    landelijk = ['Escape Game The Hunt',
                 'Outdoor escape room in ruim 25 Nederlandse steden: puzzels oplossen en ontsnappen aan de Hunters. Vanaf 8 personen.',
                 'Escape Game The Hunt is een outdoor escape room die je speelt in de binnenstad van ruim 25 Nederlandse steden, van Groningen tot Maastricht. '
                 'Geïnspireerd op tv-programma\'s als Hunted en Jachtseizoen. ' + HOE + ' ' + WIE.format(mn=P['min_people'], mx=P['max_people'])
                 + ' Ook op maat te maken voor grote groepen en evenementen.',
                 'Per stad een vaste startlocatie, zie website', URL + '/escape-game-the-hunt-locaties/',
                 PRIJS.format(b=P['base'], bm=P['base_max_people'], pp=P['per_person']), '90 minuten (met uitleg ongeveer 2 uur)',
                 f'{P["min_people"]} tot ca. {P["max_people"]} personen', URL + SITE.get('og_image', ''), '']
    wb = Workbook(); ws = wb.active; ws.title = 'Vermeldingen'
    intro = ['Teksten voor vermeldingen op uitjessites, VVV-sites en Tripadvisor. Eén regel per stad, bovenaan de landelijke versie.',
             f'Overal hetzelfde: naam "Escape Game The Hunt [stad]", telefoon {SITE["phone_display"]}, e-mail {SITE["email"]}, '
             'categorieën: outdoor escape room, citygame, teamuitje, vrijgezellenfeest, groepsuitje.',
             'De lange teksten verschillen per stad (start, straten, opening), zodat sites niet allemaal dezelfde tekst tonen. Link altijd naar de stadspagina.']
    for i, t in enumerate(intro, 1):
        ws.cell(i, 1, t).font = Font(bold=i == 1, italic=i > 1)
    for c, h in enumerate(H, 1):
        x = ws.cell(5, c, h); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='29394A')
    for r, row in enumerate([landelijk] + rows, 6):
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, v); x.alignment = Alignment(wrap_text=True, vertical='top')
    for col, w in zip('ABCDEFGHIJ', (30, 45, 90, 30, 45, 40, 20, 20, 45, 45)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'B6'
    f = os.path.join(HERE, 'vermeldingen-per-stad.xlsx'); wb.save(f)
    print(len(rows), 'steden ->', f)
    long_ = max(len(r[1]) for r in rows + [landelijk]); print('langste korte omschrijving:', long_)
    print(rows[0][2])


if __name__ == '__main__':
    main()
